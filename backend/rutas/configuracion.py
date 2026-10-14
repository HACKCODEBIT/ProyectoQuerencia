from typing import List, Optional
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.http import HttpResponse
from django.utils import timezone
from django.db.models import Sum
from ninja import Router, Schema
from ninja.errors import HttpError

from rutas.auth import AuthBearer
from modelos.models import Producto, Pedido, Factura, Usuario, Reporte, RecetaInsumo

configuracion_router = Router()

def verificar_admin(request):
    if request.auth.tipo_usuario != "administrador":
        raise HttpError(403, "Se requieren permisos de administrador.")

class MetricasResponse(Schema):
    total_ventas: float
    total_pedidos: int
    facturas_count: int
    productos_insumos: int
    stock_bajo_count: int
    platos_menu: int
    usuarios_count: int

class ReporteItem(Schema):
    idReporte: int
    fechaGeneracion: str
    tipo: str
    GeneradoPor: str

@configuracion_router.get("/metricas", response=MetricasResponse, auth=AuthBearer())
def obtener_metricas(request):
    verificar_admin(request)
    
    ventas_sum = Factura.objects.aggregate(total=Sum('total'))['total'] or 0
    total_pedidos = Pedido.objects.count()
    facturas_count = Factura.objects.count()
    productos_insumos = Producto.objects.filter(es_insumo=True).count()
    stock_bajo_count = Producto.objects.filter(es_insumo=True, stock__lt=10).count()
    platos_menu = Producto.objects.filter(es_insumo=False).count()
    usuarios_count = Usuario.objects.count()
    
    return {
        "total_ventas": float(ventas_sum),
        "total_pedidos": total_pedidos,
        "facturas_count": facturas_count,
        "productos_insumos": productos_insumos,
        "stock_bajo_count": stock_bajo_count,
        "platos_menu": platos_menu,
        "usuarios_count": usuarios_count
    }

@configuracion_router.get("/historial-reportes", response=List[ReporteItem], auth=AuthBearer())
def listar_historial_reportes(request):
    verificar_admin(request)
    reportes = Reporte.objects.all().order_by('-fechaGeneracion', '-idReporte')[:25]
    return [
        {
            "idReporte": r.idReporte,
            "fechaGeneracion": str(r.fechaGeneracion),
            "tipo": r.tipo,
            "GeneradoPor": r.GeneradoPor
        }
        for r in reportes
    ]

@configuracion_router.get("/reporte/general-excel", auth=AuthBearer())
def exportar_reporte_general(request):
    verificar_admin(request)
    
    # Registrar en el historial de reportes
    nombre_usuario = getattr(request.auth, 'nombre', 'Administrador')
    Reporte.objects.create(
        fechaGeneracion=timezone.now().date(),
        tipo="Reporte Ejecutivo Consolidado (Excel)",
        GeneradoPor=nombre_usuario
    )
    
    wb = openpyxl.Workbook()
    header_fill = PatternFill(start_color="5240CF", end_color="5240CF", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    title_font = Font(bold=True, size=14, color="1E293B")
    
    # ------------------ HOJA 1: RESUMEN EJECUTIVO ------------------
    ws_resumen = wb.active
    ws_resumen.title = "Resumen Ejecutivo"
    ws_resumen.append(["REPORTE EJECUTIVO CONSOLIDADO - PROYECTO LA QUERENCIA"])
    ws_resumen.append([f"Fecha de generación: {timezone.now().strftime('%Y-%m-%d %H:%M:%S')}"])
    ws_resumen.append([f"Generado por: {nombre_usuario}"])
    ws_resumen.append([])
    
    ventas_sum = Factura.objects.aggregate(total=Sum('total'))['total'] or 0
    kpis = [
        ["Indicador / Métrica", "Valor"],
        ["Total Ventas Facturadas ($)", float(ventas_sum)],
        ["Total de Pedidos Registrados", Pedido.objects.count()],
        ["Total Facturas Emitidas", Factura.objects.count()],
        ["Total Insumos en Inventario", Producto.objects.filter(es_insumo=True).count()],
        ["Insumos con Stock Bajo (<10)", Producto.objects.filter(es_insumo=True, stock__lt=10).count()],
        ["Total Platos en el Menú", Producto.objects.filter(es_insumo=False).count()],
        ["Usuarios Registrados en Sistema", Usuario.objects.count()]
    ]
    for row in kpis:
        ws_resumen.append(row)
        
    for col in ws_resumen.iter_cols(min_row=5, max_row=5+len(kpis)-1, min_col=1, max_col=2):
        for cell in col:
            if cell.row == 5:
                cell.fill = header_fill
                cell.font = header_font
            else:
                cell.alignment = Alignment(horizontal="left")
    
    # ------------------ HOJA 2: INVENTARIO DE INSUMOS ------------------
    ws_inv = wb.create_sheet(title="Inventario Insumos")
    ws_inv.append(["DNI / Código", "Nombre del Insumo", "Categoría", "Stock Actual", "Estado Stock"])
    for cell in ws_inv[1]:
        cell.fill = header_fill
        cell.font = header_font
        
    for p in Producto.objects.filter(es_insumo=True).order_by("nombre"):
        cat_nom = p.categoria.nombre if p.categoria else "Sin categoría"
        estado = "⚠️ CRÍTICO" if p.stock <= 5 else ("⚠️ BAJO" if p.stock < 10 else "NORMAL")
        ws_inv.append([p.dni_producto, p.nombre, cat_nom, p.stock, estado])
        
    # ------------------ HOJA 3: MENÚ Y RECETAS ------------------
    ws_menu = wb.create_sheet(title="Platos del Menú")
    ws_menu.append(["Código", "Nombre del Plato", "Categoría", "Precio Venta ($)", "Stock Disp.", "Insumos Requeridos"])
    for cell in ws_menu[1]:
        cell.fill = header_fill
        cell.font = header_font
        
    for p in Producto.objects.filter(es_insumo=False).order_by("nombre"):
        cat_nom = p.categoria.nombre if p.categoria else "Sin categoría"
        ing_list = [f"{rec.insumo.nombre} (x{rec.cantidad})" for rec in p.ingredientes_receta.all()]
        ing_str = ", ".join(ing_list) if ing_list else "Sin insumos registrados"
        ws_menu.append([p.dni_producto, p.nombre, cat_nom, float(p.precio), p.stock, ing_str])
        
    # ------------------ HOJA 4: FACTURACIÓN Y VENTAS ------------------
    ws_fac = wb.create_sheet(title="Ventas y Facturación")
    ws_fac.append(["# Factura", "Fecha", "Cliente / Razón Social", "Método de Pago", "Subtotal", "IVA", "Total Pagado"])
    for cell in ws_fac[1]:
        cell.fill = header_fill
        cell.font = header_font
        
    for f in Factura.objects.all().order_by("-fecha", "-idFactura"):
        num = f.numero_factura or f"FAC-{f.idFactura}"
        cliente = f.razon_social or "Consumidor Final"
        ws_fac.append([num, str(f.fecha), cliente, f.metodo_pago or "Efectivo", float(f.subtotal), float(f.iva), float(f.total)])
        
    # Ajustar anchos automáticos en todas las hojas
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    fecha_str = datetime.now().strftime("%Y%m%d_%H%M")
    response['Content-Disposition'] = f'attachment; filename="reporte_general_querencia_{fecha_str}.xlsx"'
    wb.save(response)
    return response
